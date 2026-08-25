"""Final audit and gap-fix pass for the user-edited direction 1B DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "提交pdf初稿-方向1B优化版.docx"
OUTPUT = ROOT / "提交pdf初稿-方向1B查漏补缺版.docx"
DIAGRAM = ROOT / "tmp" / "direction1b-architecture.png"


def remove_paragraph(p):
    el = p._element
    el.getparent().remove(el)


def insert_after(paragraph, text=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    result = Paragraph(new_p, paragraph._parent)
    if text is not None:
        result.add_run(text)
    return result


def immediate_next_paragraph_text(paragraph):
    """Return the text of the next XML sibling when it is a paragraph."""
    sibling = paragraph._p.getnext()
    if sibling is None or sibling.tag != qn("w:p"):
        return ""
    from docx.text.paragraph import Paragraph
    return Paragraph(sibling, paragraph._parent).text


def set_text(p, text, style=None):
    p.text = text
    if style:
        p.style = style
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08


def make_diagram(path: Path):
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 1800, 520
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = next((p for p in [Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/simhei.ttf")] if p.exists()), None)
    font = ImageFont.truetype(str(font_path), 26) if font_path else ImageFont.load_default()
    small = ImageFont.truetype(str(font_path), 21) if font_path else ImageFont.load_default()
    labels = [
        ("研究目标", "问题/对象/指标"),
        ("任务规划", "干预/对照\n固定seed"),
        ("执行前校验", "schema/隐私\n停止条件"),
        ("虚拟实验", "BKT+ACT-R\n/L-Model"),
        ("数据分析", "g/CI/稳健性\n异常"),
        ("人工反馈", "专家判断/边界"),
        ("第二轮计划", "变更审计/再执行"),
    ]
    x0, box_w, gap, y = 35, 220, 35, 150
    for i, (title, sub) in enumerate(labels):
        x = x0 + i * (box_w + gap)
        fill = (231, 240, 252) if i not in (5, 6) else (247, 239, 218)
        draw.rounded_rectangle((x, y, x + box_w, y + 150), radius=12, fill=fill, outline=(65, 105, 160), width=3)
        tw = draw.textbbox((0, 0), title, font=font)[2]
        draw.text((x + (box_w - tw) / 2, y + 25), title, fill=(20, 45, 80), font=font)
        lines = sub.splitlines()
        for j, line in enumerate(lines):
            tw = draw.textbbox((0, 0), line, font=small)[2]
            draw.text((x + (box_w - tw) / 2, y + 78 + j * 28), line, fill=(55, 55, 55), font=small)
        if i < len(labels) - 1:
            ax = x + box_w + 8
            ay = y + 75
            draw.line((ax, ay, ax + gap - 14, ay), fill=(75, 75, 75), width=4)
            draw.polygon([(ax + gap - 14, ay), (ax + gap - 25, ay - 8), (ax + gap - 25, ay + 8)], fill=(75, 75, 75))
    draw.text((50, 55), "方向1B：科学实验任务规划与反馈迭代闭环", fill=(20, 45, 80), font=font)
    img.save(path)


def set_cell(cell, text):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.size = Pt(9)


def main():
    make_diagram(DIAGRAM)
    doc = Document(str(SOURCE))

    # Remove all plain empty paragraphs while preserving both user images.
    for p in list(doc.paragraphs)[::-1]:
        has_drawing = "w:drawing" in p._p.xml
        if not p.text.strip() and not has_drawing:
            remove_paragraph(p)

    # Normalize every P1-P20 title that survived Word's editing pass.
    p_heading_re = re.compile(r"^P(?:[1-9]|1[0-9]|20)(?=\uFF5C|\||\s|$)")
    for p in doc.paragraphs:
        if p_heading_re.match(p.text.strip()):
            p.style = "Heading 1"
            for run in p.runs:
                run.bold = True
                run.font.size = None
        elif p.style.name == "Heading 1":
            # Only P1-P20 are section titles; body text must not inherit a
            # heading style from an edited template paragraph.
            p.style = "Normal"
        if p.text.startswith("人工审阅由"):
            set_text(p, "人工审阅由研究者完成：系统提供统计摘要、异常原因和变更建议，研究者决定是否补充证据、重跑或进入真人实验。", "Normal")

    # Add missing factual paragraphs under the sections that previously had
    # only a heading/table, keeping the actual run boundaries explicit.
    p13 = next((p for p in doc.paragraphs if p.text.strip() == "案例的实际初始条件"), None)
    if p13 is not None:
        if "case-20260821-100" not in immediate_next_paragraph_text(p13):
            p = insert_after(p13, "case-20260821-100：seed=20260821，100名学生、20名教师、100名家长，90天仿真，5个干预臂与1个对照组；数据为合成画像，运行模式为offline deterministic fallback。")
            p.style = "Normal"
    p15 = next((p for p in doc.paragraphs if (p.text.strip().startswith("P15") and "实验执行" in p.text) or p.text.strip() == "第一轮实验执行"), None)
    if p15 is not None:
        if "第一轮执行状态" not in immediate_next_paragraph_text(p15):
            p = insert_after(p15, "第一轮执行状态：代表性案例已完成虚拟执行并生成日轨迹、事件日志、5臂效应量和95% CI；本轮未接入真人样本或真实实验仪器。")
            p.style = "Normal"
    p17 = next((p for p in doc.paragraphs if p.text.strip() == "第二轮结论"), None)
    if p17 is not None:
        existing = p17._p.getnext()
        if existing is None or "未重新执行" not in "".join(existing.itertext()):
            p = insert_after(p17, "说明：第二轮完成的是工程与文档层面的反馈修订，尚未重新运行真人实验；因此第二轮结论只表示计划质量和审计边界改善。")
            p.style = "Normal"

    # Add a real architecture diagram after the P6 architecture paragraph.
    arch = next((p for p in doc.paragraphs if p.text.startswith("真实架构：")), None)
    if arch is not None and not any("direction1b-architecture" in (r.text or "") for r in arch.runs):
        p = insert_after(arch)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(DIAGRAM), width=Inches(6.5))

    # Replace generic template filler rows with meaningful status rows.
    replacements = {
        (6, 5): ["反馈与停止", "专家意见/异常", "下一轮变更清单", "需人工确认"],
        (12, 4): ["停止与归档", "CI/异常/证据不足", "暂停升级并保留日志", "下一轮计划或项目终止", "研究者复核后决定", "工程审计记录", "需人工确认"],
        (13, 4): ["真人验证准备", "真实样本与伦理方案", "尚未采集真人数据", "待伦理、授权与样本条件", "不纳入当前效果结论", "后续研究者负责"],
        (16, 5): ["停止与归档", "CI/异常/证据不足", "暂停升级", "风险控制", "保留失败轮次"],
    }
    for (ti, ri), values in replacements.items():
        if ti < len(doc.tables) and ri < len(doc.tables[ti].rows):
            for ci, value in enumerate(values):
                if ci < len(doc.tables[ti].rows[ri].cells):
                    set_cell(doc.tables[ti].rows[ri].cells[ci], value)

    # Make the data/evidence table explicitly name the literature baseline
    # module, while keeping the fixed-seed artifact in the same row.
    if len(doc.tables) > 2 and len(doc.tables[2].rows) > 4:
        set_cell(doc.tables[2].rows[4].cells[0], "文献基线/固定seed")
        set_cell(doc.tables[2].rows[4].cells[1], "literature_reference.py + case-20260821-100")

    # Uniform body formatting and no accidental blank table cells.
    for table in doc.tables:
        for ri, row in enumerate(table.rows):
            for cell in row.cells:
                if not cell.text.strip():
                    set_cell(cell, "当前未执行（边界已说明）")
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.left_indent = Pt(0)
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT

    doc.save(str(OUTPUT))
    print(f"wrote {OUTPUT}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
