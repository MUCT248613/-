from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(r"D:\Desktop\-")
SOURCE = ROOT / "提交pdf初稿-最终交付版.docx"
OUTPUT = ROOT / "提交pdf初稿-无提示交付版.docx"


PROMPT_MARKERS = (
    "证据管理原则：",
    "评价原则：",
    "125 题测试状态：",
    "案例迭代原则：",
    "逐题记录原则：",
    "对照范围说明：",
)


def main() -> None:
    doc = Document(str(SOURCE))
    removed = []
    for p in list(doc.paragraphs):
        if any(marker in p.text for marker in PROMPT_MARKERS):
            removed.append(p.text)
            p._element.getparent().remove(p._element)
    doc.core_properties.title = "虚拟学生试验台 v6.0 技术方案（无提示交付版）"
    doc.core_properties.comments = "已删除模板提示段落；正式章节保留对应事实、边界和未完成项。"
    doc.save(str(OUTPUT))
    check = Document(str(OUTPUT))
    print(f"wrote {OUTPUT}; removed={len(removed)}; tables={len(check.tables)}; images={len(check.inline_shapes)}")
    print("remaining_markers", [p.text for p in check.paragraphs if any(m in p.text for m in PROMPT_MARKERS)])


if __name__ == "__main__":
    main()
